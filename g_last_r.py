#!/usr/bin/env python3
"""
Generate laporan HTML dari semua data di transactions.db.

Pakai:
  python g_last_r.py              -> tulis ke report.html
  python g_last_r.py hasil.html   -> tulis ke file yang ditentuin
  python g_last_r.py hasil.xlsx   -> versi Excel
"""

import os
import sys
import math
import sqlite3
from datetime import datetime
from html import escape as _escape

# selalu transactions.db di folder yang sama dengan script, dari mana pun dijalanin
DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "transactions.db")

def escape(value):
    """escape() versi aman, handle None/angka tanpa crash."""
    return _escape(str(value) if value is not None else "")

PALETTE = [
    "#4E79A7", "#F28E2B", "#E15759", "#76B7B2", "#59A14F", "#EDC948",
    "#B07AA1", "#FF9DA7", "#9C755F", "#BAB0AC", "#86BCB6", "#D37295",
    "#FABFD2", "#B6992D", "#499894",
]

CREATED_FORMATS = (
    "%Y%m%d%H%M%S",
    "%Y-%m-%d %H:%M:%S.%fZ",
    "%Y-%m-%d %H:%M:%S%z",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S.%fZ",
    "%Y-%m-%dT%H:%M:%S",
)


def parse_created(value):
    value = str(value).strip()
    for fmt in CREATED_FORMATS:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    raise ValueError(f"Tidak bisa parse tanggal: {value!r}")


def fetch_report_data():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute("SELECT * FROM transactions ORDER BY created").fetchall()
    conn.close()
    result = []
    for r in rows:
        result.append({
            "id": r["id"],
            "type": r["type"],
            "category": r["category"],
            "quantity": r["quantity"],
            "amount": r["amount"],
            "description": r["description"],
            "created_at": parse_created(r["created"]),
            "is_outlier": r["is_outlier"],
            # DB lama belum punya kolom account
            "account": r["account"] if "account" in r.keys() else None,
        })
    return result


# ---------- Chart / color helpers (pure SVG, no dependency) ----------

def color_scale(value, min_v, max_v, low_hex, mid_hex, high_hex):
    if max_v == min_v:
        ratio = 0.5
    else:
        ratio = (value - min_v) / (max_v - min_v)
    if ratio < 0.5:
        r = int(low_hex[0:2], 16) + (int(mid_hex[0:2], 16) - int(low_hex[0:2], 16)) * ratio * 2
        g = int(low_hex[2:4], 16) + (int(mid_hex[2:4], 16) - int(low_hex[2:4], 16)) * ratio * 2
        b = int(low_hex[4:6], 16) + (int(mid_hex[4:6], 16) - int(low_hex[4:6], 16)) * ratio * 2
    else:
        r = int(mid_hex[0:2], 16) + (int(high_hex[0:2], 16) - int(mid_hex[0:2], 16)) * (ratio - 0.5) * 2
        g = int(mid_hex[2:4], 16) + (int(high_hex[2:4], 16) - int(mid_hex[2:4], 16)) * (ratio - 0.5) * 2
        b = int(mid_hex[4:6], 16) + (int(high_hex[4:6], 16) - int(mid_hex[4:6], 16)) * (ratio - 0.5) * 2
    return f"#{int(r):02x}{int(g):02x}{int(b):02x}"


def pie_chart_svg(items):
    total = sum(v for _, v in items)
    if total == 0:
        return ""
    start_angle = 0
    parts = [f'<svg width="300" height="300" xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 300">']
    for idx, (label, value) in enumerate(items):
        angle = (value / total) * 360
        start_rad = math.radians(start_angle)
        end_rad = math.radians(start_angle + angle)
        large_arc = 1 if angle > 180 else 0
        x1 = 150 + 100 * math.cos(start_rad)
        y1 = 150 + 100 * math.sin(start_rad)
        x2 = 150 + 100 * math.cos(end_rad)
        y2 = 150 + 100 * math.sin(end_rad)
        color = PALETTE[idx % len(PALETTE)]
        parts.append(
            f'<path d="M 150 150 L {x1} {y1} A 100 100 0 {large_arc} 1 {x2} {y2} Z" '
            f'fill="{color}" stroke="white" stroke-width="2"/>'
        )
        mid_angle = math.radians(start_angle + angle / 2)
        label_x = 150 + 65 * math.cos(mid_angle)
        label_y = 150 + 65 * math.sin(mid_angle)
        parts.append(
            f'<text x="{label_x}" y="{label_y}" text-anchor="middle" '
            f'font-size="11" fill="white" font-weight="bold">{escape(f"{value:,.0f}")}</text>'
        )
        start_angle += angle
    parts.append("</svg>")
    legend = []
    for idx, (label, value) in enumerate(items):
        pct = (value / total * 100) if total > 0 else 0
        color = PALETTE[idx % len(PALETTE)]
        legend.append(
            f'<div class="legend-item"><span class="swatch" style="background:{color}"></span>'
            f'{escape(label)} &ndash; Rp{value:,.0f} ({pct:.1f}%)</div>'
        )
    return '<div class="chart-wrap"><div>' + "".join(parts) + '</div><div>' + "".join(legend) + "</div></div>"


def line_chart_svg(months, series):
    if not months:
        return ""
    width, height = 700, 300
    pad = 40
    plot_width = width - pad * 2
    plot_height = height - pad * 2
    max_val = max((max(vals) if vals else 0) for vals in series.values()) or 1
    x_pos = lambda i: pad + (i / (len(months) - 1 if len(months) > 1 else 1)) * plot_width
    y_pos = lambda v: pad + plot_height - (v / max_val * plot_height)
    parts = [
        f'<svg width="{width}" height="{height}" xmlns="http://www.w3.org/2000/svg">',
        f'<rect width="{width}" height="{height}" fill="white"/>',
        f'<line x1="{pad}" y1="{height - pad}" x2="{width - pad}" y2="{height - pad}" stroke="#ccc" stroke-width="1"/>',
        f'<line x1="{pad}" y1="{pad}" x2="{pad}" y2="{height - pad}" stroke="#ccc" stroke-width="1"/>',
    ]

    colors = {"Income": "#2E7D32", "Expense": "#C62828"}
    legend = []
    for label, values in series.items():
        color = colors.get(label, "#333")
        points = " ".join(f"{x_pos(i):.1f},{y_pos(v):.1f}" for i, v in enumerate(values))
        parts.append(f'<polyline points="{points}" fill="none" stroke="{color}" stroke-width="2"/>')
        for i, v in enumerate(values):
            parts.append(f'<circle cx="{x_pos(i):.1f}" cy="{y_pos(v):.1f}" r="3" fill="{color}"/>')
        legend.append(f'<span class="swatch" style="background:{color}"></span>{escape(label)}')

    for i, m in enumerate(months):
        x, y = x_pos(i), height - pad + 14
        parts.append(
            f'<text x="{x:.1f}" y="{y:.1f}" font-size="10" text-anchor="end" '
            f'transform="rotate(-45 {x:.1f} {y:.1f})">{escape(m)}</text>'
        )

    parts.append("</svg>")
    legend_html = '<div class="legend">' + " &nbsp;&nbsp; ".join(legend) + "</div>"
    return '<div class="chart-wrap"><div>' + "".join(parts) + "</div>" + legend_html + "</div>"


STYLE = """
<style>
  * { box-sizing: border-box; }
  body { 
    font-family: Arial, Helvetica, sans-serif; 
    margin: 0; 
    padding: 16px; 
    color: #222; 
    background: #fff;
  }
  h1 { margin: 0 0 4px; font-size: 24px; }
  h2 { margin: 24px 0 12px; font-size: 18px; }
  .muted { color: #777; margin-top: 4px; font-size: 13px; }
  
  /* Table responsive */
  .table-wrapper { width: 100%; overflow-x: auto; margin: 12px 0 28px; }
  table { 
    border-collapse: collapse; 
    width: 100%; 
    font-size: 13px; 
    min-width: 800px;
  }
  th, td { 
    border: 1px solid #ccc; 
    padding: 8px 6px; 
    text-align: right;
  }
  th { 
    background: #ADD8E6; 
    font-weight: bold;
    position: sticky;
    top: 0;
  }
  td:first-child, th:first-child { text-align: left; }
  tr:nth-child(even) td { background: #fafafa; }
  tr:hover td { background: #f0f0f0 !important; }
  
  /* Ratio columns styling */
  .ratio-value { font-weight: 600; }
  .ratio-good { color: #2E7D32; }
  .ratio-warn { color: #F57C00; }
  .ratio-bad { color: #C62828; }
  
  /* Chart responsive */
  .chart-wrap { 
    display: flex; 
    align-items: center; 
    gap: 20px; 
    flex-wrap: wrap; 
    margin-bottom: 8px; 
    max-width: 100%; 
    overflow-x: auto; 
  }
  .chart-wrap svg { 
    max-width: 100%; 
    height: auto; 
  }
  .legend { 
    font-size: 12px; 
    display: flex;
    flex-wrap: wrap;
    gap: 16px;
  }
  .legend-item { 
    margin-bottom: 4px; 
    white-space: nowrap;
  }
  .swatch { 
    display: inline-block; 
    width: 10px; 
    height: 10px; 
    margin-right: 6px; 
    border-radius: 2px; 
    vertical-align: middle;
  }
  
  /* Mobile responsif */
  @media (max-width: 768px) {
    body { padding: 12px; }
    h1 { font-size: 20px; }
    h2 { font-size: 16px; margin: 16px 0 8px; }
    table { font-size: 12px; min-width: 600px; }
    th, td { padding: 6px 4px; }
    .table-wrapper { margin: 8px 0 20px; }
    .chart-wrap { gap: 12px; }
    .legend { font-size: 11px; gap: 12px; }
  }
  
  @media (max-width: 480px) {
    body { padding: 8px; }
    h1 { font-size: 18px; }
    h2 { font-size: 14px; margin: 12px 0 6px; }
    table { font-size: 11px; min-width: 550px; }
    th, td { padding: 5px 3px; }
    .ratio-value { font-weight: 600; }
    .legend { font-size: 10px; gap: 8px; }
  }
</style>
"""


def generate_html_report(data, out_path):
    month_data = {}
    monthly_totals = {}
    for row in data:
        month_str = row["created_at"].strftime("%Y%m")
        month_data.setdefault(month_str, []).append(row)
        totals = monthly_totals.setdefault(
            month_str, {"income": 0, "expense_clean": 0, "expense_outlier": 0}
        )
        if row["type"] == "expense":
            if row["is_outlier"]:
                totals["expense_outlier"] += row["amount"] or 0
            else:
                totals["expense_clean"] += row["amount"] or 0
        else:
            totals["income"] += row["amount"] or 0

    months_sorted = sorted(monthly_totals.keys())
    months_sorted_desc = sorted(monthly_totals.keys(), reverse=True)
    incomes = [monthly_totals[m]["income"] for m in months_sorted]
    expenses = [monthly_totals[m]["expense_clean"] + monthly_totals[m]["expense_outlier"] for m in months_sorted]

    # Cuma bulan sampe bulan berjalan yang di-color grade. Bulan yang
    # kecatet di masa depan (misal salah input tanggal) dibiarin polos.
    current_month = datetime.now().strftime("%Y%m")
    graded_months = [m for m in months_sorted if m <= current_month]

    income_vals = [monthly_totals[m]["income"] for m in graded_months] or [0]
    exp_clean_vals = [monthly_totals[m]["expense_clean"] for m in graded_months] or [0]
    exp_out_vals = [monthly_totals[m]["expense_outlier"] for m in graded_months] or [0]
    income_min, income_max = min(income_vals), max(income_vals)
    exp_clean_min, exp_clean_max = min(exp_clean_vals), max(exp_clean_vals)
    exp_out_min, exp_out_max = min(exp_out_vals), max(exp_out_vals)

    html = ["<html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width, initial-scale=1'><title>Laporan Transaksi</title>", STYLE, "</head><body>"]
    html.append("<h1>Laporan Transaksi</h1>")
    html.append(f"<p class='muted'>Dibuat {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>")

    html.append("<h2>Summary</h2>")
    html.append(line_chart_svg(months_sorted, {"Income": incomes, "Expense": expenses}))
    html.append("<div class='table-wrapper'>")
    html.append(
        "<table><thead><tr><th>Month</th><th>Income</th>"
        "<th>Expense (Outlier Excluded)</th><th>Expense Outlier</th><th>Expense Ratio</th><th>Saving %</th></tr></thead><tbody>"
    )
    for m in months_sorted_desc:
        t = monthly_totals[m]
        total_income = t["income"]
        expense_clean = t["expense_clean"]
        
        # Calculate ratios
        if total_income > 0:
            expense_ratio = (expense_clean / total_income) * 100
            saving_pct = ((total_income - expense_clean) / total_income) * 100
        else:
            expense_ratio = 0
            saving_pct = 0
        
        # Determine color for ratios
        if expense_ratio <= 50:
            exp_ratio_class = "ratio-good"
        elif expense_ratio <= 75:
            exp_ratio_class = "ratio-warn"
        else:
            exp_ratio_class = "ratio-bad"
        
        if saving_pct >= 30:
            saving_class = "ratio-good"
        elif saving_pct >= 10:
            saving_class = "ratio-warn"
        else:
            saving_class = "ratio-bad"
        
        if m <= current_month:
            income_color = color_scale(t["income"], income_min, income_max, "F8696B", "FFEB84", "63BE7B")
            exp_clean_color = color_scale(t["expense_clean"], exp_clean_min, exp_clean_max, "63BE7B", "FFEB84", "F8696B")
            exp_out_color = color_scale(t["expense_outlier"], exp_out_min, exp_out_max, "63BE7B", "FFEB84", "F8696B")
            income_style = f" style='background:{income_color}'"
            exp_clean_style = f" style='background:{exp_clean_color}'"
            exp_out_style = f" style='background:{exp_out_color}'"
        else:
            # bulan di masa depan: ga di-grade
            income_style = exp_clean_style = exp_out_style = ""
        html.append(
            "<tr>"
            f"<td>{escape(m)}</td>"
            f"<td{income_style}>Rp{t['income']:,.0f}</td>"
            f"<td{exp_clean_style}>Rp{t['expense_clean']:,.0f}</td>"
            f"<td{exp_out_style}>Rp{t['expense_outlier']:,.0f}</td>"
            f"<td class='ratio-value {exp_ratio_class}'>{expense_ratio:.1f}%</td>"
            f"<td class='ratio-value {saving_class}'>{saving_pct:.1f}%</td>"
            "</tr>"
        )
    html.append("</tbody></table>")
    html.append("</div>")

    for m in sorted(month_data.keys(), reverse=True):
        rows = month_data[m]
        html.append(f"<h2>{escape(m)}</h2>")

        cat_totals = {}
        for r in rows:
            if r["type"] == "expense" and not r["is_outlier"]:
                cat = r["category"] or "(tanpa kategori)"
                cat_totals[cat] = cat_totals.get(cat, 0) + (r["amount"] or 0)
        if cat_totals:
            html.append(pie_chart_svg(sorted(cat_totals.items(), key=lambda kv: -kv[1])))

        html.append("<div class='table-wrapper'>")
        html.append(
            "<table><thead><tr><th>ID</th><th>Type</th><th>Category</th><th>Qty</th>"
            "<th>Amount</th><th>Description</th><th>Created</th><th>Outlier</th></tr></thead><tbody>"
        )
        for r in sorted(rows, key=lambda x: x["created_at"], reverse=True):
            html.append(
                "<tr>"
                f"<td>{escape(r['id'])}</td>"
                f"<td>{escape(r['type'])}</td>"
                f"<td>{escape(r['category'])}</td>"
                f"<td>{r['quantity'] if r['quantity'] is not None else ''}</td>"
                f"<td>Rp{(r['amount'] or 0):,.0f}</td>"
                f"<td>{escape(r['description'])}</td>"
                f"<td>{r['created_at'].strftime('%Y-%m-%d %H:%M:%S')}</td>"
                f"<td>{'Yes' if r['is_outlier'] else ''}</td>"
                "</tr>"
            )
        html.append("</tbody></table>")
        html.append("</div>")

    html.append("</body></html>")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(html))



def generate_xlsx_report(data, out_path):
    """Versi Excel dari laporan: sheet Summary + satu sheet per bulan (terbaru duluan)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    money = '"Rp"#,##0'
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="ADD8E6")

    def write_header(ws, headers):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws.freeze_panes = "A2"

    def set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    month_data = {}
    monthly_totals = {}
    for row in data:
        month_str = row["created_at"].strftime("%Y%m")
        month_data.setdefault(month_str, []).append(row)
        totals = monthly_totals.setdefault(
            month_str, {"income": 0, "expense_clean": 0, "expense_outlier": 0}
        )
        if row["type"] == "expense":
            bucket = "expense_outlier" if row["is_outlier"] else "expense_clean"
        else:
            bucket = "income"
        totals[bucket] += row["amount"] or 0

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_header(ws, [
        "Month", "Income", "Expense (Outlier Excluded)",
        "Expense Outlier", "Expense Ratio", "Saving", "Saving %",
    ])
    for i, m in enumerate(sorted(monthly_totals, reverse=True), 2):
        t = monthly_totals[m]
        # Saving = income - expense (outlier excluded), sama dasarnya dengan Saving %.
        # Ditulis sebagai angka jadi, bukan formula: viewer yang ga ngitung formula
        # (preview Telegram, viewer HP) bakal nampilin sel formula kosong.
        income = t["income"]
        saving = income - t["expense_clean"]
        ws.append([
            m, income, t["expense_clean"], t["expense_outlier"],
            t["expense_clean"] / income if income else 0,
            saving,
            saving / income if income else 0,
        ])
        for col in "BCDF":
            ws[f"{col}{i}"].number_format = money
        for col in "EG":
            ws[f"{col}{i}"].number_format = "0.0%"
    # Sama kayak HTML: cuma bulan sampe bulan berjalan yang diwarnain. Bulan di masa
    # depan (salah input tanggal) ada di baris paling atas karena urutan desc, jadi
    # range warnanya mulai dari baris bulan berjalan/terakhir sebelumnya.
    current_month = datetime.now().strftime("%Y%m")
    months_desc = sorted(monthly_totals, reverse=True)
    future_count = sum(1 for m in months_desc if m > current_month)
    first = 2 + future_count
    last = ws.max_row
    if last >= first:
        # income makin gede makin hijau, expense makin gede makin merah; tiap kolom diskala sendiri
        good_high = ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                                   mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B")
        bad_high = ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile",
                                  mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B")
        for col in "BFG":
            ws.conditional_formatting.add(f"{col}{first}:{col}{last}", good_high)
        for col in "CDE":
            ws.conditional_formatting.add(f"{col}{first}:{col}{last}", bad_high)
        # Saving minus (expense > income) dibikin merah biar langsung keliatan
        ws.conditional_formatting.add(
            f"F{first}:F{last}",
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C62828", bold=True)),
        )
    set_widths(ws, [10, 16, 26, 16, 14, 16, 12])

    for m in sorted(month_data, reverse=True):
        ws = wb.create_sheet(m)
        write_header(ws, ["ID", "Type", "Category", "Qty", "Amount", "Description", "Created", "Outlier"])
        for r in sorted(month_data[m], key=lambda x: x["created_at"], reverse=True):
            ws.append([
                r["id"], r["type"], r["category"], r["quantity"], r["amount"] or 0,
                r["description"], r["created_at"], "Yes" if r["is_outlier"] else "",
            ])
            ws.cell(ws.max_row, 5).number_format = money
            ws.cell(ws.max_row, 7).number_format = "yyyy-mm-dd hh:mm:ss"
        set_widths(ws, [18, 10, 16, 6, 14, 40, 20, 8])

    wb.save(out_path)


def main():
    out_path = sys.argv[1] if len(sys.argv) > 1 else "report.html"
    if not os.path.exists(DB_PATH):
        raise SystemExit(f"{DB_PATH} ga ketemu")
    data = fetch_report_data()
    if not data:
        raise SystemExit("Belum ada data transaksi.")
    if out_path.lower().endswith(".xlsx"):
        generate_xlsx_report(data, out_path)
    else:
        generate_html_report(data, out_path)
    print(f"Laporan tersimpan: {out_path}")


if __name__ == "__main__":
    main()
